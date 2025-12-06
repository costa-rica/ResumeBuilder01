import docx

doc = docx.Document()
doc.add_heading('John Doe', 0)
doc.add_paragraph('Software Engineer with 5 years of experience in Python and JavaScript.')
doc.add_heading('Experience', level=1)
doc.add_paragraph('Senior Developer at Tech Corp (2020-Present)')
doc.add_paragraph('Developed web applications using FastAPI and React.')
doc.save('test_resume.docx')
print("Created test_resume.docx")
